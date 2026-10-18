import io
import pypdfium2 as pdfium
from docx import Document
from PIL import Image
from pypdf import PdfReader
from app.vision_client import read_image_text

MAX_CHARS = 40000
CHUNK_SIZE = 450
CHUNK_OVERLAP = 80
MIN_TEXT_PER_PAGE = 40
MAX_OCR_PAGES = 10
IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".webp")


def _normalize(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _jpeg(image: Image.Image) -> bytes:
    image = image.convert("RGB")
    image.thumbnail((1800, 1800))
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=80)
    return buffer.getvalue()


def _ocr(image: Image.Image, label: str) -> str:
    try:
        return read_image_text(_jpeg(image))
    except Exception as error:
        print(f"OCR error ({label}): {error}")
        return ""


def _pdf_text(data: bytes) -> tuple[str, int]:
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return text, len(reader.pages)


def _ocr_pdf(data: bytes) -> str:
    pdf = pdfium.PdfDocument(data)
    total = min(len(pdf), MAX_OCR_PAGES)
    parts = []
    for index in range(total):
        print(f"OCR page {index + 1}/{total}")
        image = pdf[index].render(scale=2).to_pil()
        parts.append(_ocr(image, f"page {index + 1}"))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            text, pages = _pdf_text(data)
            if len(text.strip()) < MIN_TEXT_PER_PAGE * max(pages, 1):
                print("PDF has no text layer, using OCR")
                text = _ocr_pdf(data)
        elif name.endswith(".docx"):
            document = Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        elif name.endswith((".txt", ".md", ".csv")):
            text = data.decode("utf-8", errors="ignore")
        elif name.endswith(IMAGE_EXTENSIONS):
            print("Image attachment, using OCR")
            text = _ocr(Image.open(io.BytesIO(data)), filename)
        else:
            return ""
    except Exception as error:
        print(f"Text extract error ({filename}): {error}")
        return ""
    return _normalize(text)[:MAX_CHARS]


def split_chunks(text: str) -> list[str]:
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + CHUNK_SIZE)
        if end < len(text):
            cut = text.rfind("\n", start + CHUNK_SIZE // 2, end)
            if cut == -1:
                cut = text.rfind(". ", start + CHUNK_SIZE // 2, end)
            if cut != -1:
                end = cut + 1
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = end - CHUNK_OVERLAP
    return chunks